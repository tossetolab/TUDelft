#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Step2の分析CSVを1つのExcelブックに集約 (Step 3)"""
import argparse, csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

HDR = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
BASE = Font(name="Arial", size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def read_csv(path):
    with open(path, encoding="utf-8-sig") as f:
        return list(csv.reader(f))

def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = HDR; cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

def dump(ws, rows, widths=None):
    for r in rows:
        ws.append(r)
    style_header(ws, len(rows[0]))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BASE; cell.border = BORDER
    if widths:
        for i, wv in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wv
    ws.freeze_panes = "A2"

def num_cols(ws, cols, minrow=2):
    for row in ws.iter_rows(min_row=minrow):
        for ci in cols:
            row[ci].number_format = "#,##0"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--work", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    w = args.work
    wb = Workbook()

    # --- 概要シート ---
    ws = wb.active; ws.title = "概要"
    ws["A1"] = "PLATEAU技術レポート 重要語分析（試作）"
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    notes = [
        "",
        "【対象】PLATEAU技術検証レポート（試作: 5冊）",
        "【処理】PDFテキスト抽出 → fugashi(MeCab)形態素解析 → 頻度/TF-IDF/共起/文書スコア",
        "",
        "【シート構成】",
        "  ・文書一覧＆スコア : 各レポートの語数・特徴の強さ(TF-IDF上位合計)・代表語",
        "  ・全体頻出語       : コーパス全体の出現回数と出現文書数（DF）",
        "  ・特徴語(文書別)   : 各レポートのTF-IDF上位語（そのレポートを特徴づける語）",
        "  ・共起ペア         : 同一文中で共起する語ペア（概念のつながり）",
        "",
        "【指標の意味】",
        "  TF-IDF … そのレポートで相対的に特徴的な語ほど高い（全レポート共通の語は低い）",
        "  DF(出現文書数) … 語がいくつのレポートに現れたか。全冊共通=一般語の目安",
        "",
        "※本ブックは5冊の試作結果です。同じスクリプトで120冊へそのまま拡張できます。",
    ]
    for i, t in enumerate(notes, 3):
        ws.cell(i, 1, t).font = BASE
    ws.column_dimensions["A"].width = 90

    # --- 文書一覧＆スコア ---
    rows = read_csv(os.path.join(w, "doc_scores.csv"))
    rows[0] = ["ファイル", "タイトル", "延べ語数", "異なり語数", "特徴強度(TF-IDF上位和)", "代表語トップ3"]
    ws = wb.create_sheet("文書一覧＆スコア")
    dump(ws, rows, [30, 46, 12, 12, 20, 30])
    num_cols(ws, [2, 3])

    # --- 全体頻出語 (+棒グラフ) ---
    rows = read_csv(os.path.join(w, "corpus_freq.csv"))
    rows[0] = ["語", "総出現回数", "出現文書数(DF)"]
    rows = [rows[0]] + rows[1:201]  # 上位200語
    ws = wb.create_sheet("全体頻出語")
    dump(ws, rows, [24, 14, 16])
    num_cols(ws, [1, 2])
    chart = BarChart(); chart.title = "頻出語 上位20"; chart.type = "bar"
    chart.height = 12; chart.width = 20
    data = Reference(ws, min_col=2, min_row=1, max_row=21)
    cats = Reference(ws, min_col=1, min_row=2, max_row=21)
    chart.add_data(data, titles_from_data=True); chart.set_categories(cats)
    ws.add_chart(chart, "E2")

    # --- 特徴語(文書別) ---
    rows = read_csv(os.path.join(w, "feature_terms.csv"))
    rows[0] = ["ファイル", "タイトル", "順位", "特徴語", "TF-IDF", "そのレポート内出現回数"]
    ws = wb.create_sheet("特徴語(文書別)")
    dump(ws, rows, [30, 40, 6, 20, 10, 20])

    # --- 共起ペア ---
    rows = read_csv(os.path.join(w, "cooccurrence.csv"))
    rows[0] = ["語1", "語2", "共起回数(同一文)"]
    ws = wb.create_sheet("共起ペア")
    dump(ws, rows, [22, 22, 16])
    num_cols(ws, [2])

    wb.save(args.out)
    print("saved", args.out)

if __name__ == "__main__":
    main()
